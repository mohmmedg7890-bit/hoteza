import io,json,uuid,zipfile,tempfile
from pathlib import Path
from datetime import timedelta
from unittest.mock import patch
from django.test import TestCase,Client,TransactionTestCase,override_settings
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import connection,close_old_connections
from django.utils import timezone
from .models import Property,Membership,Unit,Guest,Reservation,Allocation,Charge,Payment,Invoice,Audit,Contract,Expense
from .services import perform,RuleError,balance,report,money

class PMSCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner=get_user_model().objects.create_user('owner',password='Test-owner!926')
        cls.staff=get_user_model().objects.create_user('staff',password='Test-staff!926')
        cls.other=get_user_model().objects.create_user('other',password='Test-other!926')
        cls.p=Property.objects.create(name='منشأة أ',tax_rate=15)
        cls.q=Property.objects.create(name='منشأة ب',tax_rate=15)
        Membership.objects.create(property=cls.p,user=cls.owner,role='owner')
        Membership.objects.create(property=cls.p,user=cls.staff,role='reception')
        Membership.objects.create(property=cls.q,user=cls.other,role='owner')
        cls.u=Unit.objects.create(property=cls.p,number='101',kind='استوديو',capacity=2)
        cls.v=Unit.objects.create(property=cls.p,number='102',kind='غرفة وصالة',capacity=3)
        cls.w=Unit.objects.create(property=cls.q,number='101',kind='استوديو')
        cls.g=Guest.objects.create(property=cls.p,name='نزيل تجريبي',phone='0500000000',email='guest@example.com',region='عسير',city='أبها',street='شارع الاختبار')
        cls.h=Guest.objects.create(property=cls.q,name='سري في منشأة ب')
    def setUp(self):self.today=timezone.localdate();self.client.force_login(self.owner)
    def op(self,action,d,user=None,key=None):
        actor=user or self.owner;key=key or str(uuid.uuid4())
        if action=='checkout':
            from .checkout import preview
            cache=getattr(self,'checkout_payloads',{});self.checkout_payloads=cache
            if key in cache:d=cache[key]
            else:
                plan=preview(actor,self.p.pk,d['id']);d={**d,'settlement_ack':True,'settlement_digest':plan['digest'],'settlement_mode':'debt' if d.get('allow_debt') else 'none','debt_reason':'اعتماد مديونية للاختبار فقط'};cache[key]=d
        return perform(actor,self.p.pk,action,d,key)
    def book(self,**kw):
        d={'guest':self.g.pk,'unit':self.u.pk,'arrival':str(self.today),'departure':str(self.today+timedelta(days=3)),'rate':'150','persons':2};d.update(kw)
        return Reservation.objects.get(pk=self.op('reservation_create',d)['id'])
    def api(self,action,d,key=None):return self.client.post('/api/action/'+action+'/',data=json.dumps(d),content_type='application/json',HTTP_IDEMPOTENCY_KEY=key or str(uuid.uuid4()))
    def test_entry_forms_are_document_pages(self):
        response=self.client.get('/app/')
        self.assertContains(response,'<section id="modal" class="entry-page"')
        self.assertNotContains(response,'<dialog id="modal"')
        self.assertContains(response,'/static/entry-pages.js')
    def test_settlement_register_is_tenant_scoped_and_read_only(self):
        r=self.book();other=Reservation.objects.create(property=self.q,guest=self.h,arrival=self.today,departure=self.today+timedelta(days=1),rate=100)
        for prop,res in [(self.p,r),(self.q,other)]:Audit.objects.create(property=prop,user=self.owner,action='checkout_settlement',details={'reservation':res.pk,'after':{'total':100},'mode':'none'})
        count=Audit.objects.count();result=self.client.get('/api/vouchers/settlements/').json()
        self.assertEqual(result['total'],1);self.assertEqual(result['rows'][0]['reservation'],r.pk)
        self.assertEqual(Audit.objects.count(),count)
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get('/api/vouchers/settlements/').json()['total'],1)
    def test_expense_document_is_printable_and_tenant_scoped(self):
        x=Expense.objects.create(property=self.p,amount=100,method='نقد',beneficiary='مستفيد اختبار',description='سند صرف اختبار',user=self.owner)
        self.assertContains(self.client.get(f'/document/expense/{x.pk}/'),'سند صرف')
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(f'/document/expense/{x.pk}/').status_code,400)
    def test_strict_guest_fields_and_saudi_normalization(self):
        valid={'first_name':'محمد','last_name':'أحمد','gender':'male','birth_date':'1990-01-01','nationality':'السعودية','customer_type':'citizen','id_type':'national_id','id_number':'1000000000','phone':'٠٥٠٠٠٠٠٠٠٠','email':'guest@example.com','region':'عسير','city':'أبها','street':'شارع الاختبار'}
        response=self.api('guest_save',valid);self.assertEqual(response.status_code,200)
        self.assertEqual(Guest.objects.get(pk=response.json()['id']).phone,'+966500000000')
        for key,value in [('first_name',''),('last_name','123'),('phone','+9660500000000'),('phone','051'),('email','invalid'),('region',''),('city',''),('street','')]:
            with self.subTest(key=key,value=value):self.assertEqual(self.api('guest_save',{**valid,key:value}).status_code,400)
    def profile_data(self,**extra):
        return {'first_name':'محمد','second_name':'أحمد','last_name':'علي','first_name_en':'Mohammed','phone':'0500000000','email':'guest@example.com','region':'عسير','city':'أبها','street':'عنوان الاختبار','gender':'male','birth_date':'1990-01-01','nationality':'السعودية','customer_type':'citizen','id_type':'national_id','id_number':'1000000000',**extra}
    def test_profile_identity_birth_and_optional_fields(self):
        data=self.profile_data(workplace='شركة الاختبار',vehicle_plate='أ ب ج 123')
        response=self.api('guest_save',data);self.assertEqual(response.status_code,200)
        guest=Guest.objects.get(pk=response.json()['id']);self.assertEqual(guest.name,'محمد أحمد علي')
        self.assertEqual(guest.profile['first_name_en'],'Mohammed');self.assertTrue(guest.profile['birth_hijri'].startswith('1410-'))
        state=self.client.get('/api/state/').json();saved=next(g for g in state['guests'] if g['id']==guest.pk)
        self.assertEqual(saved['id_number'],'1000000000');self.assertNotIn('photo',saved)
        for extra in [{'gender':''},{'birth_date':'2099-01-01'},{'birth_date':'1900-01-01'},{'nationality':''},{'customer_type':'invalid'},{'id_type':'iqama'},{'id_number':'123'},{'birth_hijri':'1400-01-01'},{'first_name_en':'محمد'}]:
            with self.subTest(extra=extra):self.assertEqual(self.api('guest_save',self.profile_data(**extra)).status_code,400)
        self.assertEqual(self.client.get('/api/birth-date/?date=1990-01-01').json()['birth_hijri'],guest.profile['birth_hijri'])
        self.assertEqual(self.client.get('/api/birth-date/?date=2099-01-01').status_code,400)
    def test_photo_persist_scoped_and_removable(self):
        import base64
        from PIL import Image
        image=io.BytesIO();Image.new('RGB',(400,300),'blue').save(image,format='JPEG')
        data='data:image/jpeg;base64,'+base64.b64encode(image.getvalue()).decode()
        response=self.api('guest_save',self.profile_data(photo_data=data,show_photo=True));self.assertEqual(response.status_code,200)
        pk=response.json()['id'];photo=self.client.get(f'/api/guest-photo/{pk}/')
        self.assertEqual(photo.status_code,200);self.assertEqual(photo['Content-Type'],'image/jpeg');self.assertIn('no-store',photo['Cache-Control'])
        self.client.force_login(self.other);self.assertEqual(self.client.get(f'/api/guest-photo/{pk}/').status_code,404)
        self.client.force_login(self.owner)
        self.assertEqual(self.api('guest_save',self.profile_data(id=pk,photo_data='')).status_code,200)
        self.assertEqual(self.client.get(f'/api/guest-photo/{pk}/').status_code,404)
        self.assertEqual(self.api('guest_save',self.profile_data(photo_data='data:image/jpeg;base64,bad')).status_code,400)
    def test_report_export_selected_view(self):
        self.book();params=f'?start={self.today}&end={self.today+timedelta(days=3)}'
        for view,heading,excluded in [('units','الوحدة','المصدر'),('sources','المصدر','الوحدة'),('daily','التاريخ','الوحدة')]:
            response=self.client.get('/export/reports/json/'+params+'&view='+view);self.assertEqual(response.status_code,200)
            rows=response.json()['rows'];self.assertTrue(any(row[0]==heading for row in rows));self.assertFalse(any(row[0]==excluded for row in rows))

    def test_expected_times_validate_and_guest_update_rolls_back(self):
        bad=self.api('reservation_create',{'guest':self.g.pk,'unit':self.u.pk,'arrival':str(self.today),'departure':str(self.today+timedelta(days=1)),'arrival_time':'29:00'})
        self.assertEqual(bad.status_code,400);self.assertEqual(Reservation.objects.count(),0)
        r=self.book(arrival_time='16:30',departure_time='11:00');self.assertEqual(r.arrival_time.strftime('%H:%M'),'16:30')
        details={'name':'اسم مختلف','phone':'0500000000','email':'new@example.com','region':'عسير','city':'أبها','street':'عنوان جديد'}
        response=self.api('reservation_create',{'guest':self.g.pk,'unit':self.u.pk,'arrival':str(self.today),'departure':str(self.today+timedelta(days=1)),'guest_details':details})
        self.assertEqual(response.status_code,400);self.g.refresh_from_db();self.assertEqual(self.g.name,'نزيل تجريبي')
    def test_legacy_guest_requires_completion_before_new_booking(self):
        self.g.city='';self.g.save()
        response=self.api('reservation_create',{'guest':self.g.pk,'unit':self.u.pk,'arrival':str(self.today),'departure':str(self.today+timedelta(days=1))})
        self.assertEqual(response.status_code,400);self.assertEqual(Reservation.objects.count(),0)
    def test_monthly_multi_unit_pricing_and_atomic_conflict(self):
        from datetime import date
        from .operations import add_months
        with patch('pms.services.timezone.localdate',return_value=date(2027,1,31)):
            r=self.book(arrival='2027-01-31',departure='2027-02-28',rental_type='monthly',units=[{'unit':self.u.pk,'rate':'2800'},{'unit':self.v.pk,'rate':'1400'}],persons=4)
            self.assertEqual(balance(r)['net'],420000);self.assertEqual(balance(r)['total'],483000)
            self.assertEqual(r.allocations.count(),2);self.assertEqual(r.contracts.count(),1)
            self.assertEqual(sum(c.net for c in r.charges.filter(unit=self.u)),280000)
            with self.assertRaises(RuleError):self.book(arrival='2027-01-31',departure='2027-02-28',units=[{'unit':self.v.pk,'rate':100},{'unit':self.u.pk,'rate':100}])
            self.assertEqual(Reservation.objects.count(),1)
            self.op('checkin',{'id':r.pk});self.op('extend',{'id':r.pk,'departure':'2027-03-31','rate':4200})
            self.assertEqual(balance(r)['net'],840000)
        self.assertEqual(add_months(date(2027,1,31),1),date(2027,2,28))
    def test_discount_expense_and_contract_snapshot(self):
        r=self.book();self.op('discount',{'id':r.pk,'discount_type':'percent','discount_value':10})
        self.assertEqual(balance(r)['total'],46575)
        before=balance(r);result=self.op('expense',{'id':r.pk,'amount':100,'method':'نقد','beneficiary':'مورد تجريبي','description':'شراء مستلزمات'})
        self.assertEqual(balance(r),before);self.assertEqual(Expense.objects.get(pk=result['expense']).amount,10000)
        self.op('invoice',{'id':r.pk});self.assertEqual(r.invoices.first().total,46575)
        with self.assertRaises(RuleError):self.op('discount',{'id':r.pk,'discount_type':'fixed','discount_value':10})
        self.op('contract_issue',{'id':r.pk});c=r.contracts.order_by('-id').first();original=c.snapshot
        self.g.name='اسم آخر';self.g.save();c.refresh_from_db();self.assertEqual(c.snapshot,original)
        response=self.client.get(f'/document/contract/{c.pk}/');self.assertEqual(response.status_code,200);self.assertContains(response,'مسودة غير مكتملة')
        self.client.force_login(self.other);self.assertEqual(self.client.get(f'/document/contract/{c.pk}/').status_code,400)
    def test_contract_terms_share_expiration_and_tamper(self):
        r=self.book()
        with self.assertRaises(RuleError):self.op('contract_share',{'id':r.pk})
        terms=[f'شرط اختباري كامل رقم {i}' for i in range(14)]
        self.op('terms_save',{'terms':terms});self.op('contract_issue',{'id':r.pk})
        link=self.op('contract_share',{'id':r.pk})['url'];anon=Client()
        self.assertEqual(anon.get(link).status_code,200);self.assertContains(anon.get(link),terms[-1])
        self.assertEqual(anon.get(link[:-2]+'x/').status_code,404)
        with patch('django.core.signing.time.time',return_value=timezone.now().timestamp()+73*3600):self.assertEqual(anon.get(link).status_code,404)
        self.assertEqual(self.api('terms_save',{'terms':terms[:13]}).status_code,400)
    def test_multi_unit_checkout_and_auto_extension(self):
        r=self.book(units=[{'unit':self.u.pk,'rate':100},{'unit':self.v.pk,'rate':100}],persons=3,auto_extend=True)
        self.op('checkin',{'id':r.pk})
        with patch('pms.services.timezone.localdate',return_value=r.departure):self.op('auto_extend',{'id':r.pk})
        r.refresh_from_db();self.assertEqual(r.departure,self.today+timedelta(days=4));self.assertEqual(r.allocations.filter(end=r.departure).count(),2)
        self.op('invoice',{'id':r.pk});self.op('payment',{'id':r.pk,'amount':920,'method':'شبكة'})
        with patch('pms.services.timezone.localdate',return_value=r.departure):self.op('checkout',{'id':r.pk})
        self.u.refresh_from_db();self.v.refresh_from_db();self.assertEqual([self.u.cleaning,self.v.cleaning],['dirty','dirty'])

    def test_same_day_checkout_retains_one_night_and_reverses_discount(self):
        r=self.book(discount_type='percent',discount_value=10)
        self.op('checkin',{'id':r.pk})
        self.op('early_departure',{'id':r.pk,'departure':str(self.today+timedelta(days=1))})
        self.assertEqual(balance(r)['total'],15525)
        self.op('invoice',{'id':r.pk});self.op('payment',{'id':r.pk,'amount':'155.25','method':'نقد'})
        self.op('checkout',{'id':r.pk,'allow_same_day':True});r.refresh_from_db();self.assertEqual(r.status,'checked_out')

    def test_add_remove_units_preserves_guest_balance_and_billed_lines(self):
        r=self.book(persons=2);initial=balance(r)
        self.op('unit_add',{'id':r.pk,'unit':self.v.pk,'rate':100});self.assertEqual(balance(r)['net'],75000)
        allocation=r.allocations.get(unit=self.v);self.op('unit_remove',{'id':r.pk,'allocation':allocation.pk})
        self.assertEqual(balance(r),initial);self.assertEqual(r.allocations.count(),1)
        with self.assertRaises(RuleError):self.op('unit_remove',{'id':r.pk,'allocation':r.allocations.first().pk})
        self.op('unit_add',{'id':r.pk,'unit':self.v.pk,'rate':100});self.op('invoice',{'id':r.pk})
        with self.assertRaises(RuleError):self.op('unit_remove',{'id':r.pk,'allocation':r.allocations.get(unit=self.v).pk})

    def test_booking_amounts_and_tax(self):
        r=self.book();self.assertEqual(balance(r),{'net':45000,'tax':6750,'total':51750,'paid':0,'balance':51750});self.assertEqual(r.allocations.count(),1)
    def test_overlap_rejected_without_partial_guest_charge(self):
        self.book()
        with self.assertRaises(RuleError):self.book()
        self.assertEqual(Reservation.objects.count(),1);self.assertEqual(Charge.objects.count(),3)
    def test_adjacent_booking_allowed(self):
        r=self.book();self.book(arrival=str(r.departure),departure=str(r.departure+timedelta(days=1)));self.assertEqual(Reservation.objects.count(),2)
    def test_capacity_and_invalid_dates(self):
        for extra in [{'persons':3},{'departure':str(self.today)},{'arrival':'nonsense'},{'rate':'NaN'},{'rate':'1.001'},{'rate':'-1'}]:
            with self.subTest(extra=extra),self.assertRaises(RuleError):self.book(**extra)
        self.assertEqual(Reservation.objects.count(),0)
    def test_idempotent_booking_and_mismatched_reuse(self):
        key=str(uuid.uuid4());data={'guest':self.g.pk,'unit':self.u.pk,'arrival':str(self.today),'departure':str(self.today+timedelta(days=1)),'rate':150}
        a=self.op('reservation_create',data,key=key);b=self.op('reservation_create',data,key=key);self.assertEqual(a,b);self.assertEqual(Reservation.objects.count(),1)
        with self.assertRaises(RuleError):self.op('reservation_create',{**data,'rate':200},key=key)
    def test_cross_tenant_fk_injection(self):
        for fields in [{'unit':self.w.pk},{'guest':self.h.pk}]:
            with self.subTest(fields=fields),self.assertRaises(RuleError):self.book(**fields)
        self.assertFalse(Reservation.objects.exists())
    def test_cross_tenant_mutation_and_read(self):
        self.assertEqual(self.api('unit_save',{'id':self.w.pk,'number':'x','kind':'x'}).status_code,400)
        self.assertEqual(self.api('select_property',{'id':self.q.pk}).status_code,400)
        s=self.client.get('/api/state/').json();self.assertEqual(len(s['units']),2);self.assertNotIn(self.h.name,json.dumps(s,ensure_ascii=False))
    def test_reception_cannot_change_settings_users_or_refund(self):
        self.client.force_login(self.staff)
        for action in ['property_save','unit_save','user_add','refund','adjustment','early_departure','subscription']:
            with self.subTest(action=action):self.assertEqual(self.api(action,{}).status_code,400)
        self.assertEqual(self.client.get('/api/state/').json()['users'],[])
    def test_csrf_and_anonymous_protection(self):
        anon=Client();self.assertEqual(anon.get('/api/state/').status_code,401)
        strict=Client(enforce_csrf_checks=True);strict.force_login(self.owner)
        self.assertEqual(strict.post('/api/action/guest_save/',data='{}',content_type='application/json').status_code,403)
        self.assertEqual(anon.get('/export/reservations/csv/').status_code,401)
    def test_unavailable_cleaning_blocks_checkin(self):
        r=self.book();self.u.cleaning='dirty';self.u.save()
        with self.assertRaises(RuleError):self.op('checkin',{'id':r.pk})
        self.op('cleaning',{'id':self.u.pk,'cleaning':'ready'});self.op('checkin',{'id':r.pk});r.refresh_from_db();self.assertEqual(r.status,'checked_in')
    def test_extend_conflict_is_atomic(self):
        r=self.book();self.book(arrival=str(r.departure),departure=str(r.departure+timedelta(days=1)))
        before=Charge.objects.count()
        with self.assertRaises(RuleError):self.op('extend',{'id':r.pk,'departure':str(r.departure+timedelta(days=2)),'rate':150})
        r.refresh_from_db();self.assertEqual(r.departure,self.today+timedelta(days=3));self.assertEqual(Charge.objects.count(),before)
    def test_extend_adds_only_new_nights(self):
        r=self.book();self.op('extend',{'id':r.pk,'departure':str(r.departure+timedelta(days=2)),'rate':200});self.assertEqual(balance(r)['net'],85000);self.assertEqual(r.charges.count(),5)
    def test_move_frees_old_and_records_audit(self):
        r=self.book();self.op('checkin',{'id':r.pk});self.op('move',{'id':r.pk,'unit':self.v.pk});self.u.refresh_from_db()
        self.assertEqual(self.u.cleaning,'dirty');self.assertEqual(r.allocations.last().unit_id,self.v.pk)
        self.assertEqual(Audit.objects.filter(action='move').last().details['result']['from_unit'],self.u.pk)
    def test_move_conflict_does_not_alter_original(self):
        r=self.book();self.book(unit=self.v.pk);self.op('checkin',{'id':r.pk})
        with self.assertRaises(RuleError):self.op('move',{'id':r.pk,'unit':self.v.pk})
        self.assertEqual(r.allocations.last().unit_id,self.u.pk)
    def test_maintenance_rejects_overlap_and_blocks_booking(self):
        self.op('block',{'unit':self.u.pk,'start':str(self.today),'end':str(self.today+timedelta(days=2)),'reason':'صيانة'})
        with self.assertRaises(RuleError):self.book()
        self.assertEqual(report(self.p,self.today,self.today+timedelta(days=1))['capacity'],1)
    def test_block_cannot_replace_existing_reservation(self):
        self.book()
        with self.assertRaises(RuleError):self.op('block',{'unit':self.u.pk,'start':str(self.today),'end':str(self.today+timedelta(days=1)),'reason':'صيانة'})
    def test_cancel_releases_availability_and_preserves_record(self):
        r=self.book();self.op('cancel',{'id':r.pk});self.book();r.refresh_from_db();self.assertEqual(r.status,'cancelled');self.assertEqual(r.charges.count(),0)
    def test_edit_recalculates_and_cannot_override_receipted_booking(self):
        r=self.book();d={'id':r.pk,'unit':self.v.pk,'arrival':str(self.today),'departure':str(self.today+timedelta(days=2)),'rate':200,'persons':2}
        self.op('reservation_edit',d);self.assertEqual(balance(r)['net'],40000)
        self.op('payment',{'id':r.pk,'amount':50,'method':'نقد'})
        with self.assertRaises(RuleError):self.op('reservation_edit',d)
    def test_receipt_refund_and_immutable_invoice_snapshot(self):
        r=self.book();i=self.op('invoice',{'id':r.pk,'buyer':'شركة الاختبار','buyer_tax':'123456789012345'})['invoice']
        original=Invoice.objects.get(pk=i).snapshot
        self.op('property_save',{'name':'اسم معدل','tax_rate':0});self.g.name='اسم معدل';self.g.save();self.assertEqual(Invoice.objects.get(pk=i).snapshot,original)
        self.op('payment',{'id':r.pk,'amount':600,'method':'شبكة'});self.op('refund',{'id':r.pk,'amount':'82.50','method':'شبكة'})
        self.assertEqual(balance(r)['balance'],0);self.assertEqual(r.payments.count(),2)
        with self.assertRaises(RuleError):self.op('refund',{'id':r.pk,'amount':1,'method':'نقد'})
    def test_invoice_only_unbilled_and_credit_limits(self):
        r=self.book();inv=self.op('invoice',{'id':r.pk})['invoice']
        with self.assertRaises(RuleError):self.op('invoice',{'id':r.pk})
        self.op('adjustment',{'id':r.pk,'invoice':inv,'kind':'credit','amount':'17.25','reason':'خصم'})
        self.assertEqual(balance(r)['total'],50025)
        with self.assertRaises(RuleError):self.op('adjustment',{'id':r.pk,'invoice':inv,'kind':'credit','amount':600,'reason':'خطأ'})
    def test_checkout_requires_settlement_and_financial_document(self):
        r=self.book();self.op('checkin',{'id':r.pk});future=self.today+timedelta(days=3)
        with patch('pms.services.timezone.localdate',return_value=future):
            with self.assertRaises(RuleError):self.op('checkout',{'id':r.pk})
            self.op('payment',{'id':r.pk,'amount':'517.50','method':'نقد'})
            key=str(uuid.uuid4());result=self.op('checkout',{'id':r.pk},key=key)
            self.assertEqual(result,self.op('checkout',{'id':r.pk},key=key))
            self.assertEqual(len(result['invoices']),1);self.assertEqual(r.invoices.count(),1)
            self.assertFalse(r.charges.filter(invoice__isnull=True).exists())
            self.assertEqual(Invoice.objects.get(pk=result['invoices'][0]).total,51750)
        r.refresh_from_db();self.u.refresh_from_db();self.assertEqual(r.status,'checked_out');self.assertEqual(self.u.cleaning,'dirty')
    def test_cash_movement_separates_revenue_payments_refunds_and_expenses(self):
        r=self.book();self.op('payment',{'id':r.pk,'amount':'600','method':'نقد'})
        self.op('refund',{'id':r.pk,'amount':'20','method':'نقد'})
        self.op('expense',{'amount':'30','method':'نقد','beneficiary':'مورد تجريبي','description':'مصروف اختبار'})
        rep=report(self.p,self.today,self.today+timedelta(days=1))
        self.assertEqual(rep['paid'],58000);self.assertEqual(rep['expenses'],3000);self.assertEqual(rep['cash_movement'],55000)
        self.assertEqual(rep['net'],15000);self.assertEqual(balance(r)['paid'],58000)
        self.assertEqual(report(self.q,self.today,self.today+timedelta(days=1))['expenses'],0)
        response=self.client.get(f'/export/reports/json/?start={self.today}&end={self.today+timedelta(days=1)}')
        self.assertEqual(response.status_code,200);self.assertIn('المصروفات المدفوعة',str(response.json()))

    def test_owner_can_approve_debt_reception_cannot(self):
        r=self.book();self.op('checkin',{'id':r.pk});self.op('invoice',{'id':r.pk})
        with patch('pms.services.timezone.localdate',return_value=self.today+timedelta(days=3)):
            with self.assertRaises(RuleError):self.op('checkout',{'id':r.pk,'allow_debt':True},user=self.staff)
            self.op('checkout',{'id':r.pk,'allow_debt':True})
        self.assertEqual(balance(r)['balance'],51750)
    def test_early_departure_credits_unused_billed_nights(self):
        r=self.book();self.op('checkin',{'id':r.pk});self.op('invoice',{'id':r.pk});self.op('payment',{'id':r.pk,'amount':'517.50','method':'نقد'})
        tomorrow=self.today+timedelta(days=1)
        with patch('pms.services.timezone.localdate',return_value=tomorrow):
            self.op('early_departure',{'id':r.pk,'departure':str(tomorrow)})
            self.assertEqual(balance(r)['total'],17250);self.assertEqual(balance(r)['balance'],-34500)
            self.op('refund',{'id':r.pk,'amount':345,'method':'نقد'});self.op('checkout',{'id':r.pk})
        self.assertEqual(r.invoices.filter(kind='credit').count(),1)
        rep=report(self.p,tomorrow,tomorrow+timedelta(days=1));self.assertEqual(rep['net'],0);self.assertEqual(rep['adr'],0)
    def test_overdue_inhouse_guest_still_blocks_unit(self):
        r=self.book();self.op('checkin',{'id':r.pk})
        with patch('pms.services.timezone.localdate',return_value=r.departure):
            with self.assertRaises(RuleError):self.book(arrival=str(r.departure),departure=str(r.departure+timedelta(days=1)))
    def test_report_splits_months_and_excludes_cancelled(self):
        d=self.today.replace(day=1)+timedelta(days=32);d=d.replace(day=1)
        self.book(arrival=str(d-timedelta(days=1)),departure=str(d+timedelta(days=2)))
        rep=report(self.p,d,d+timedelta(days=2));self.assertEqual(rep['sold'],2);self.assertEqual(rep['net'],30000);self.assertEqual(rep['capacity'],4);self.assertEqual(rep['occupancy'],50)
        self.assertEqual(sum(x['net'] for x in rep['by_source']),rep['net']);self.assertEqual(sum(x['net'] for x in rep['by_unit']),rep['net'])
        self.assertEqual(len(rep['by_unit']),2);self.assertEqual(rep['by_unit'][1]['sold'],0)
    def test_subscription_readonly_after_expiry(self):
        self.p.subscription_until=self.today-timedelta(days=1);self.p.save()
        self.assertEqual(self.client.get('/api/state/').status_code,200)
        with self.assertRaises(RuleError):self.book()
    def test_audit_does_not_store_passwords(self):
        self.op('user_add',{'username':'newuser','password':'N3w!complex$Test482','role':'reception'})
        self.assertNotIn('N3w!complex',str(list(Audit.objects.values())))
    def test_export_files_are_valid_and_private(self):
        from openpyxl import load_workbook
        from docx import Document
        from PIL import Image
        r=self.book();inv=self.op('invoice',{'id':r.pk})['invoice']
        for fmt in ['csv','json','xlsx','docx','pdf','png']:
            with self.subTest(fmt=fmt):
                resp=self.client.get(f'/export/invoice-{inv}/{fmt}/');self.assertEqual(resp.status_code,200);raw=resp.content;self.assertIn('no-store',resp['Cache-Control'])
                if fmt=='pdf':self.assertTrue(raw.startswith(b'%PDF'))
                if fmt=='xlsx':self.assertEqual(load_workbook(io.BytesIO(raw)).active['A1'].value,'المنشأة')
                if fmt=='docx':self.assertTrue(Document(io.BytesIO(raw)).tables)
                if fmt=='png':Image.open(io.BytesIO(raw)).verify()
                if fmt=='json':self.assertEqual(json.loads(raw)['property'],self.p.name)
        self.client.force_login(self.other);self.assertEqual(self.client.get(f'/export/invoice-{inv}/pdf/').status_code,400)
    def test_csv_formula_injection_is_neutralized(self):
        self.g.name='=1+1';self.g.save();raw=self.client.get('/export/guests/csv/').content.decode('utf-8-sig');self.assertIn("'=1+1",raw)
    def test_backup_includes_financials_and_refuses_overwrite(self):
        self.book()
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'backup.json';call_command('backup_data',str(p),stdout=io.StringIO());data=json.loads(p.read_text());self.assertTrue(any(x['model']=='pms.charge' for x in data));self.assertEqual(p.stat().st_mode & 0o777,0o600)
            with self.assertRaises(FileExistsError):call_command('backup_data',str(p))
            with self.assertRaises(CommandError):call_command('restore_data',str(p),confirm_empty_database=True)
    def test_login_logout_setup_and_password_pages(self):
        anon=Client();self.assertEqual(anon.get('/login/').status_code,200);self.assertEqual(anon.get('/password-reset/').status_code,200)
        self.assertEqual(self.client.get('/app/').status_code,200);self.assertEqual(self.client.get('/logout/').status_code,405)
        self.assertEqual(self.client.post('/logout/').status_code,302)
    @override_settings(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
    def test_password_recovery_sends_valid_link(self):
        from django.core import mail
        from django.utils.http import urlsafe_base64_encode
        from django.utils.encoding import force_bytes
        from django.contrib.auth.tokens import default_token_generator
        self.owner.email='owner@example.test';self.owner.save()
        response=Client().post('/password-reset/',{'email':self.owner.email})
        self.assertEqual(response.status_code,302);self.assertEqual(len(mail.outbox),1)
        uid=urlsafe_base64_encode(force_bytes(self.owner.pk));token=default_token_generator.make_token(self.owner)
        self.assertIn(f'/reset/{uid}/',mail.outbox[0].body)
        client=Client();response=client.get(f'/reset/{uid}/{token}/');self.assertEqual(response.status_code,302)
        response=client.post(response.url,{'new_password1':'New-Reset!8265*word','new_password2':'New-Reset!8265*word'})
        self.assertEqual(response.status_code,302);self.owner.refresh_from_db();self.assertTrue(self.owner.check_password('New-Reset!8265*word'))
    def test_change_password_keeps_current_session(self):
        response=self.api('profile',{'old_password':'Test-owner!926','password':'New-Profile!8265*word'})
        self.assertEqual(response.status_code,200);self.assertEqual(self.client.get('/api/state/').status_code,200)

class ConcurrencyCase(TransactionTestCase):
    def test_two_workers_cannot_double_book(self):
        # Run on a file-backed SQLite database or PostgreSQL. In-memory SQLite has
        # different lock behavior and is deliberately not claimed as production proof.
        if connection.vendor=='sqlite' and 'memory' in str(connection.settings_dict['NAME']):self.skipTest('Requires file-backed SQLite or PostgreSQL')
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        owner=get_user_model().objects.create_user('race-owner')
        p=Property.objects.create(name='Race');Membership.objects.create(user=owner,property=p,role='owner')
        u=Unit.objects.create(property=p,number='101',kind='studio');g=Guest.objects.create(property=p,name='Race guest',phone='0500000000',email='race@example.com',region='عسير',city='أبها',street='شارع الاختبار');today=timezone.localdate();barrier=Barrier(2)
        def attempt(_):
            close_old_connections();barrier.wait()
            try:
                perform(owner,p.pk,'reservation_create',{'guest':g.pk,'unit':u.pk,'arrival':str(today),'departure':str(today+timedelta(days=1)),'rate':150},str(uuid.uuid4()));return 'created'
            except RuleError:return 'rejected'
            finally:connection.close()
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(attempt,range(2)))
        self.assertEqual(sorted(results),['created','rejected']);self.assertEqual(Reservation.objects.count(),1)

class BootstrapCase(TestCase):
    def test_demo_bootstrap_keeps_password_and_data_on_restart(self):
        import os
        password = 'Trial-Strong!Pass_9274'
        out = io.StringIO()
        with patch.dict(os.environ, {'INITIAL_ADMIN_USERNAME':'trial-owner','INITIAL_ADMIN_PASSWORD':password,'BOOTSTRAP_DEMO':'1'}):
            call_command('bootstrap_owner', stdout=out)
        user = get_user_model().objects.get(username='trial-owner')
        self.assertTrue(user.check_password(password))
        self.assertEqual(Property.objects.count(), 2)
        self.assertEqual(Unit.objects.count(), 30)
        self.assertNotIn(password, out.getvalue())
        with patch.dict(os.environ, {'INITIAL_ADMIN_USERNAME':'trial-owner','INITIAL_ADMIN_PASSWORD':'Different-Strong!Pass_273','BOOTSTRAP_DEMO':'1'}):
            call_command('bootstrap_owner', stdout=out)
        user.refresh_from_db()
        self.assertTrue(user.check_password(password))
        self.assertEqual(Unit.objects.count(), 30)

    def test_bootstrap_rejects_weak_secret_without_creating_account(self):
        import os
        with patch.dict(os.environ, {'INITIAL_ADMIN_USERNAME':'owner','INITIAL_ADMIN_PASSWORD':'123','BOOTSTRAP_DEMO':'1'}):
            with self.assertRaises(CommandError):
                call_command('bootstrap_owner', stdout=io.StringIO())
        self.assertEqual(get_user_model().objects.count(), 0)
