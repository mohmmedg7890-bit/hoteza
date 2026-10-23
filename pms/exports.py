import csv, io, json
from datetime import timedelta
from decimal import Decimal
from django.conf import settings
from django.http import HttpResponse
from django.utils import timezone
from .models import Expense,Contract,Reservation, Invoice, Payment, Guest
from .services import RuleError, balance, report, day

def cash(n):return f'{Decimal(n)/100:,.2f}'

def document_data(prop,kind,pk):
    if kind=='contract':
        c=Contract.objects.filter(pk=pk,reservation__property=prop).first()
        if not c:raise RuleError('العقد غير موجود')
        s=c.snapshot;p=s['property'];g=s['guest']
        rows=[['رقم العقد',f'CTR-{prop.pk}-{c.pk:06}'],['المنشأة',p['name']],['العنوان',p['address']],['السجل التجاري',p.get('commercial_registration','')],['الرقم الضريبي',p['tax_number']],['الترخيص السياحي',p.get('tourism_license','')],['الهاتف',p.get('phone','')],['الجوال',p.get('mobile','')],['الرمز البريدي',p.get('postal_code','')],['مرجع الحجز',s['reference']],['النزيل',g['name']],['الجوال',g['phone']],['نوع الهوية',g.get('id_type','')],['رقم الهوية',g.get('id_number','')],['الجنسية',g.get('nationality','')],['الدخول',s['arrival']],['الخروج',s['departure']],['نوع الإيجار','شهري' if s['rental_type']=='monthly' else 'يومي'],['المجموع الجزئي',cash(s['net'])],['الضريبة',cash(s['tax'])],['الإجمالي',cash(s['total'])]]
        rows += [['الوحدة','النوع','الطابق','من','إلى']]+[[u['number'],u['kind'],u['floor'],u['from'],u['to']] for u in s['units']]
        rows += [['الشروط والأحكام' if s['terms_complete'] else 'مسودة — الشروط الأربعة عشر لم تكتمل']]+[[str(i+1),term] for i,term in enumerate(s['terms'])]+[['توقيع العميل','________________'],['توقيع الموظف',s['employee']+' ________________']]
        return 'عقد إيجار — '+('نسخة للشروط المسجلة' if s['terms_complete'] else 'مسودة غير مكتملة'),rows
    if kind=='expense':
        e=Expense.objects.filter(pk=pk,property=prop).select_related('user').first()
        if not e:raise RuleError('سند الصرف غير موجود')
        return 'سند صرف',[['المنشأة',prop.name],['رقم السند',f'EXP-{e.pk:06}'],['المستفيد',e.beneficiary],['البيان',e.description],['المبلغ',cash(e.amount)],['طريقة الدفع',e.method],['التاريخ',timezone.localtime(e.created_at).isoformat()],['المستخدم',e.user.username],['مرجع الحجز',e.reservation_id or 'غير مرتبط'],['ملاحظة','سند صرف من المنشأة؛ لا يخفض مديونية النزيل']]
    if kind=='invoice':
        obj=Invoice.objects.filter(pk=pk,reservation__property=prop).first()
        if not obj:raise RuleError('المستند غير موجود')
        s=obj.snapshot;label={'invoice':'فاتورة تجريبية','credit':'إشعار دائن تجريبي','debit':'إشعار مدين تجريبي'}[obj.kind]
        rows=[['المنشأة',s['property']],['العنوان',s['address']],['الرقم الضريبي',s['tax_number']],['رقم المستند',f'INV-{obj.pk:06}'],['تاريخ الإصدار',timezone.localtime(obj.created_at).isoformat()],['المشتري',s['buyer']],['رقم المشتري الضريبي',s.get('buyer_tax','')],['عنوان المشتري',s.get('buyer_address','')],['الحجز',s['reservation']],['الوصول / المغادرة',s['arrival']+' / '+s['departure']],['العملة',s['currency']]]
        if obj.reference_id:rows.append(['مرجع الفاتورة الأصلية',obj.reference_id])
        rows += [['البند','التاريخ','قبل الضريبة','الضريبة','الإجمالي']]+[[x['description'],x['date'],cash(x['net']),cash(x['tax']),cash(x['net']+x['tax'])] for x in s['lines']]
        rows += [['الإجمالي',cash(obj.total)]]
        return label,rows
    if kind=='receipt':
        obj=Payment.objects.select_related('reservation__guest','user').filter(pk=pk,reservation__property=prop).first()
        if not obj:raise RuleError('السند غير موجود')
        return ('سند قبض' if obj.amount>0 else 'سند رد مبلغ'),[['المنشأة',prop.name],['رقم السند',f'PAY-{obj.pk:06}'],['الحجز',obj.reservation_id],['النزيل',obj.reservation.guest.name],['المبلغ',cash(abs(obj.amount))],['الطريقة',obj.method],['المرجع',obj.reference],['التاريخ',timezone.localtime(obj.created_at).isoformat()],['الموظف',obj.user.username]]
    if kind in ['statement','checkout']:
        r=Reservation.objects.select_related('guest').filter(pk=pk,property=prop).first()
        if not r:raise RuleError('الحجز غير موجود')
        if kind=='checkout' and r.status!='checked_out':raise RuleError('لم تسجل مغادرة هذا الحجز')
        b=balance(r)
        rows=[['المنشأة',prop.name],['الحجز',r.pk],['النزيل',r.guest.name],['الحالة',r.status],['الوصول',str(r.arrival)],['المغادرة',str(r.departure)],['المغادرة الفعلية',str(r.checked_out_at or '')],['قبل الضريبة',cash(b['net'])],['الضريبة',cash(b['tax'])],['الإجمالي',cash(b['total'])],['صافي المقبوض',cash(b['paid'])],['المتبقي',cash(b['balance'])],['الحركة','التاريخ','القيمة']]
        rows += [[c.description,str(c.service_date),cash(c.net+c.tax)] for c in r.charges.all()]
        rows += [[('قبض ' if p.amount>0 else 'رد ')+p.method,timezone.localtime(p.created_at).isoformat(),cash(p.amount)] for p in r.payments.all()]
        return ('كشف حساب النزيل' if kind=='statement' else 'ملخص المغادرة والتسوية'),rows
    raise RuleError('نوع المستند غير معروف')

def safe(v):
    s=str(v if v is not None else '')
    return "'"+s if s.lstrip().startswith(('=','+','-','@')) else s

def build_export(request,prop,dataset,fmt,role):
    title='تقرير Hoteza';rows=[]
    if '-' in dataset:
        kind,pk=dataset.rsplit('-',1)
        title,rows=document_data(prop,kind,int(pk))
    elif dataset=='reservations':
        rows=[['الحجز','النزيل','الوصول','المغادرة','الحالة','الإجمالي','المقبوض','المتبقي']]
        for r in Reservation.objects.filter(property=prop).select_related('guest'):
            b=balance(r);rows.append([r.pk,r.guest.name,str(r.arrival),str(r.departure),r.status,cash(b['total']),cash(b['paid']),cash(b['balance'])])
    elif dataset=='guests':
        rows=[['المعرف','الاسم','الجوال','البريد']]+[[g.pk,g.name,g.phone,g.email] for g in Guest.objects.filter(property=prop)]
    elif dataset=='payments':
        rows=[['السند','الحجز','المبلغ','الطريقة','المرجع','التاريخ']]+[[p.pk,p.reservation_id,cash(p.amount),p.method,p.reference,timezone.localtime(p.created_at).isoformat()] for p in Payment.objects.filter(reservation__property=prop)]
    elif dataset=='reports':
        start=day(request.GET.get('start',timezone.localdate()));end=day(request.GET.get('end',start+timedelta(days=1)));r=report(prop,start,end)
        view=request.GET.get('view','overview')
        titles={'overview':'الملخص التشغيلي','daily':'الحركة اليومية','units':'أداء الوحدات','sources':'مصادر الحجوزات'}
        if view not in titles:raise RuleError('نوع التقرير غير معروف')
        title=titles[view]+' — Hoteza'
        rows=[['المنشأة',prop.name],['التقرير',titles[view]],['بداية الفترة',str(start)],['نهاية الفترة - غير مشمولة',str(end)],['العملة',prop.currency],['الطاقة المتاحة',r['capacity']],['الليالي المباعة',r['sold']],['الإشغال %',r['occupancy']],['الإيراد قبل الضريبة',cash(r['net'])],['الضريبة',cash(r['tax'])],['صافي التحصيل',cash(r['paid'])],['المصروفات المدفوعة',cash(r['expenses'])],['صافي الحركة النقدية - ليست صافي الربح',cash(r['cash_movement'])],['ADR قبل الضريبة',cash(r['adr'])],['RevPAR قبل الضريبة',cash(r['revpar'])]]
        if view in ['overview','daily']:
            rows += [['التاريخ','الطاقة','المباع','الإيراد','الضريبة','التحصيل','المصروفات','صافي الحركة النقدية']]+[[x['date'],x['available'],x['sold'],cash(x['net']),cash(x['tax']),cash(x['paid']),cash(x['expenses']),cash(x['cash_movement'])] for x in r['rows']]
            rows += [['الإجمالي',r['capacity'],r['sold'],cash(r['net']),cash(r['tax']),cash(r['paid']),cash(r['expenses']),cash(r['cash_movement'])]]
        if view in ['overview','sources']:
            rows += [['المصدر','الليالي','الإيراد','الضريبة']]+[[x['source'],x['sold'],cash(x['net']),cash(x['tax'])] for x in r['by_source']]
        if view in ['overview','units']:
            rows += [['الوحدة','النوع','الطاقة','الليالي','الإشغال %','الإيراد','الضريبة']]+[[x['number'],x['kind'],x['capacity'],x['sold'],round(100*x['sold']/x['capacity'],2) if x['capacity'] else 0,cash(x['net']),cash(x['tax'])] for x in r['by_unit']]
    else:raise RuleError('التقرير غير معروف')
    rows=[list(map(str,r)) for r in rows]
    out=io.BytesIO()
    types={'pdf':'application/pdf','xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','png':'image/png','csv':'text/csv; charset=utf-8','json':'application/json'}
    if fmt=='json':out.write(json.dumps({'title':title,'property':prop.name,'rows':rows},ensure_ascii=False,indent=2).encode())
    elif fmt=='csv':
        s=io.StringIO();w=csv.writer(s);w.writerows([[safe(v) for v in r] for r in rows]);out.write(('\ufeff'+s.getvalue()).encode())
    elif fmt=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb=Workbook();ws=wb.active;ws.title='Hoteza';ws.sheet_view.rightToLeft=True
        for r in rows:ws.append([safe(x) for x in r])
        for c in ws[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='153759')
        for row in ws:
            for c in row:c.alignment=Alignment(horizontal='right',wrap_text=True)
        for i in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(i)].width=28
        ws.freeze_panes='A2';wb.save(out)
    elif fmt=='docx':
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc=Document();doc.add_heading(title,0);doc.add_paragraph('مستند تجريبي — غير صالح كفاتورة إلكترونية نظامية')
        table=doc.add_table(rows=0,cols=max(len(r) for r in rows));table.style='Light Shading Accent 1'
        bidi=OxmlElement('w:bidiVisual');table._tbl.tblPr.append(bidi)
        for r in rows:
            cells=table.add_row().cells
            for i,v in enumerate(r):
                cells[i].text=v
                for p in cells[i].paragraphs:
                    p.alignment=2;b=OxmlElement('w:bidi');p._p.get_or_add_pPr().append(b)
        doc.save(out)
    elif fmt=='pdf':
        import arabic_reshaper
        from bidi.algorithm import get_display
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from textwrap import wrap
        pdfmetrics.registerFont(TTFont('Arabic',settings.FONT_PATH))
        c=canvas.Canvas(out,pagesize=(842,595));c.setTitle(title)
        page=0;y=0
        def header():
            nonlocal page,y
            page+=1;c.setFont('Arabic',15);c.drawRightString(805,555,get_display(arabic_reshaper.reshape(title)));c.setFont('Arabic',9);c.drawRightString(805,534,get_display(arabic_reshaper.reshape('مستند تجريبي — غير صالح كفاتورة إلكترونية نظامية')));c.drawString(35,25,str(page));y=507
        header()
        for r in rows:
            for line in wrap('   |   '.join(r),width=125,break_long_words=True):
                if y<45:c.showPage();header()
                c.setFont('Arabic',10);c.drawRightString(805,y,get_display(arabic_reshaper.reshape(line)));y-=20
        c.save()
    elif fmt=='png':
        from PIL import Image, ImageDraw, ImageFont
        from textwrap import wrap
        lines=[title,'مستند تجريبي — غير صالح كفاتورة إلكترونية نظامية']
        for r in rows:lines.extend(wrap(' | '.join(r),80))
        if len(lines)>400:raise RuleError('الصورة طويلة؛ استخدم PDF أو Excel لهذه الفترة')
        img=Image.new('RGB',(1500,max(300,100+45*len(lines))),'white');draw=ImageDraw.Draw(img);font=ImageFont.truetype(settings.FONT_PATH,24)
        for i,line in enumerate(lines):draw.text((1450,45+i*45),line,font=font,fill='#153759',anchor='ra',direction='rtl')
        img.save(out,format='PNG')
    else:raise RuleError('صيغة غير مدعومة')
    response=HttpResponse(out.getvalue(),content_type=types[fmt]);response['Content-Disposition']=f'attachment; filename="hoteza-{dataset}.{fmt}"';response['Cache-Control']='private, no-store';return response
